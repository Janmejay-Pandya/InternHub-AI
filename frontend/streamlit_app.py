import streamlit as st
import requests
from docx import Document
from io import BytesIO

st.set_page_config(
    page_title="InternHub AI",
    layout="centered"
)

st.title("🎯 InternHub AI Internship Matcher")
st.write("AI-powered internship matching and resume generation using LLMs")

def extract_resume(text: str):
    """
    Extract resume text using delimiters OR fallback headings
    """

    # Primary: delimiter-based extraction
    start = "=== RESUME (START) ==="
    end = "=== RESUME (END) ==="

    if start in text and end in text:
        return text.split(start)[1].split(end)[0].strip()

    # Fallback: heading-based extraction
    lines = text.split("\n")
    resume_lines = []
    capture = False

    for line in lines:
        if "resume" in line.lower():
            capture = True
            continue
        if "ats" in line.lower() or "confidence score" in line.lower():
            break
        if capture:
            resume_lines.append(line)

    if resume_lines:
        return "\n".join(resume_lines).strip()

    return None



def create_resume_doc(resume_text: str, name: str):
    """
    Create a Word (.docx) file from resume text
    """
    doc = Document()
    doc.add_heading(f"{name} - Resume", level=1)

    for line in resume_text.split("\n"):
        doc.add_paragraph(line)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

with st.form("internhub_form"):
    name = st.text_input("Full Name")
    skills = st.text_area("Skills (comma-separated)")
    interests = st.text_area("Interests")
    experience = st.text_area("Experience")
    jd = st.text_area("Internship Description")

    submitted = st.form_submit_button("Analyze Data")

if submitted:
    if not skills or not jd:
        st.error("Please fill at least Skills and Internship Description")
    else:
        payload = {
            "student": {
                "name": name,
                "skills": skills,
                "interests": interests,
                "experience": experience
            },
            "internship_description": jd
        }

        with st.spinner("🔍 Analyzing with AI..."):
            try:
                response = requests.post(
                    "http://localhost:8081/analyze",
                    json=payload,
                    timeout=60
                )

                if response.status_code == 200:
                    result_text = response.json()["result"]

                    st.success("✅ Analysis Complete")
                    st.markdown(result_text)

                    resume_text = extract_resume(result_text)

                    if resume_text:
                        doc_file = create_resume_doc(resume_text, name)

                        st.download_button(
                            label="📄 Download Resume (Word)",
                            data=doc_file,
                            file_name=f"{name}_Internship_Resume.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
                    else:
                        st.warning("Resume section not found in AI response")

                else:
                    st.error("Backend error. Please try again.")

            except Exception as e:
                st.error(f"Failed to connect to backend: {e}")
